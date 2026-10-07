import pandas as pd
from io import BytesIO
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from datetime import date

CODIGO_CLIENTE = "788809"

DERIVAR = {
    # Tiroides
    "T4 -Tiroxina Total T4 (ug/dl)":                       "866",
    "T4L -Tiroxina Libre T4 L (ng/dl)":                    "867",
    "T3 (Triiodotironina) (ng/dl)":                        "878",
    "T3 Libre (Triiodotironina Libre) (pg/mL)":            "9661",
    "TSH (Tirotrofina ultrasensible) (uUl/ml)":            "865",
    "ATPO / TPO-PEROXIDASA TIROIDEO, Ac. Anti-  (UI/mL)":  "8315",
    "Anticuerpos Antiperoxidasa (Ac anti TPO) (Ul/mL)":    "8315",
    # Minerales / iones
    "Zinc sérico (mg/L)":                                  "982",
    "Calcio Sérico (mg/dL)":                               "133",
    "Calcio Urinario 24 hs":                               "136.24",
    "Magnesio Sérico (mg/dL)":                             "653",
    "Fosfatemia (mg/dL)":                                  "362",
    "Cloro Sérico (mg/dL)":                                "168",
    # Vitaminas
    "Vitamina B 12 (pg/mL)":                               "938",
    "Vitamina B6 (Piridoxina) (ng/mL)":                    "9887",
    "Ácido Fólico Vitamina B9 (ng/mL)":                    "352",
    "Vitamina D 25 OH (ng/mL)":                            "9913",
    # Metabolismo / glucosa
    "Hemoglobina Glicosilada ( HbA1c) (%)":                "1070",
    "Insulina plasmática basal (mU/L)":                    "543",
    "Insulina 120´ (mU/L)":                                "543.1",
    # Proteínas
    "Proteínas totales (g/dL)":                            "763",
    "Albúmina sérica (g/dL)":                              "15",
    # Orina / riñón
    "Microalbuminuria, al azar  ()":                       "1130.00",
    "Relación microalbuminuria/ mg creatinina (ug/mg creat.)": "1130.00",
    # Reumatología
    "Péptido Citrulinado Cíclico Ác (U/mL)":              "8284",
    "Aldolasa (U/L)":                                      "18",
    # Hepatitis / Serología
    "Hepatitis B HBc Ac ()":                               "1080",
    "Hepatitis B Antígeno de superficie ()":               "1086",
    "Hepatitis C (HCV) Ac ()":                             "1095",
    "HIV 4º Generación (Ag / Ac Combinado) ()":            "63",
    "FTA - ABS ()":                                        "371",
    "Toxoplasmosis IgG (UI/mL)":                           "9571",
    "Toxoplasmosis IgM captura (Tasa s/co)":               "9580",
    # Micológico
    "Micológico Cultivo Identificación ()":                "665",
    # Enzimas
    "CPK (Creatinfosfoquinasa) (U/l)":                     "190",
    "LDH (Láctico Dehidrogenasa) (U/L)":                   "594",
    # PCR cuantitativa
    "PCR -Proteína \"C\" Reactiva cuantitativa (mg/L)":    "761",
    # Hormonas reproductivas
    "Gonadotrofina Coriónica Beta (Beta hCG)  (mUI/mL)":   "1175",
    "FSH (Hormona Foliculoestimulante) (mUI/mL)":          "370",
    "Estradiol (E2) (pg/mL)":                              "300",
    "Testosterona (ng/mL)":                                "863.2",
    "Testosterona libre sérica (pg/mL)":                   "9375",
    # Inmunoglobulinas
    "Inmunoglobulina A (IgA) (mg/dL)":                    "537",
    "Inmunoglobulina G (IgG) (mg/dL)":                    "540",
    "Inmunoglobulina M (IgM) (mg/dL)":                    "541",
    # Marcadores tumorales
    "CA 125 (U/mL)":                                       "1115",
    # Próstata
    "PSA -Antígeno Prostático Específico Total (ng/mL)":   "1000",
    # Perfil de hierro
    "Ferremia (ug/dL)":                                    "343",
    "Ferritina (ng/mL)":                                   "5230",
    "Transferrina (ug/dL)":                                "875",
    "Transferrina Índice de Saturación (%)":               "837",
    # Coagulación
    " Fibrinógeno":                                        "345",
    " DIMERO D":                                           "4418",
    # Hormonas adicionales
    " Prolactina":                                         "759",
    " LH (Hormona Luteinizante)":                          "612",
    " Progesterona sérica":                                "758",
    " ACTH- Adrenocorticotrofina":                         "6",
    " 17 Hidroxi Progesterona sérica":                     "8580",
    # Cortisol
    " Cortisol plasmático basal":                          "189",
    " Cortisol Vespertino":                                "189.1",
    " Cortisol post Dexametasona":                         "189.Dexa",
    " Cortisol Urinario":                                  "4008",
    " Cortisol Libre Urinario en 24 hs":                   "4008",
    " Cortisol saliva, 23:00hs":                           "4012",
    # PTH y marcadores tumorales
    " PTH- PARATHORMONA":                                  "739",
    " CEA (Antígeno Carcinoembrionario)":                  "144",
    " CA 19-9":                                            "1125",
    " Alfafetoproteína":                                   "20",
    # Hueso
    " Beta Cross Laps":                                    "3025",
    " Osteocalcina":                                       "7939",
    # Metabolismo adicional
    " Homocisteína basal":                                 "6452",
    " Homocisteína post carga de Metionina":               "6452.1",
    # Serología viral
    " Rubeola IgG":                                        "1145",
    " Rubeola IgM":                                        "1150",
    " Anticuerpos Anti Rubeola IgG específica":            "1145",
    " Anticuerpos Anti Rubeola Ig M específica":           "1150",
    " Citomegalovirus (CMV) IgG":                          "1025",
    " Citomegalovirus (CMV) IgM ":                         "1030",
    " HERPES Simplex 1 IgG":                               "6042",
    " Herpes Simplex 1 IgM":                               "6050",
    " HERPES Simplex 2 IgG":                               "6067",
    " Herpes Simplex 2 IgM":                               "6076",
    " Chagas Ac, Electroquimioluminiscencia":              "243",
    " Chagas (HAI)":                                       "242",
    " Chagas IgG IFI":                                     "243.1",
    " Hepatitis A (HAV) Ac Totales":                       "5888",
    " Hepatitis A (HAV) IgM":                              "1075",
    # Celiaquía
    " Transglutaminasa IgA":                               "9622.BF",
    " Transglutaminasa IgG":                               "9631",
    # PSA libre
    " Antígeno Prostático Específico Libre (PSA L)":       "1000.L",
    # Coprocultivo
    " Examen Microscópico Directo -Coprocultivo-":         "187",
    # Anticuerpos antifosfolípidos
    " Beta 2 GLICOPROTEÍNA, Ac. IgG Anti":                "8218",
    " Beta 2 GLICOPROTEÍNA, Ac. IgM Anti- ":              "8222",
    " Hepatitis B HBc IgM":                               "1082",
}

NOMBRE_ANALISIS = {
    "866": "T4 Total (Tiroxina Total)",
    "867": "T4L (Tiroxina Libre)",
    "878": "T3 Total (Triiodotironina)",
    "9661": "T3L (Triiodotironina Libre)",
    "865": "TSH (Tirotrofina ultrasensible)",
    "8315": "ATPO Peroxidasa Tiroidea Ac",
    "982": "Zinc sérico",
    "133": "Calcio sérico total",
    "136.24": "Calcio urinario 24 hs",
    "653": "Magnesio sérico",
    "362": "Fósforo sérico",
    "168": "Cloro (Cl) sérico",
    "938": "Vitamina B12 (Cianocobalamina)",
    "352": "Ácido Fólico sérico",
    "9913": "Vitamina D 25 OH",
    "1070": "Hemoglobina Glicosilada (A1C), HPLC",
    "543": "Insulina basal",
    "543.1": "Insulina 120' post prandial",
    "763": "Proteínas totales",
    "15": "Albúmina sérica",
    "1130.00": "Microalbuminuria, al azar",
    "8284": "Péptido Cíclico Citrulinado Ac (CCP Ac)",
    "18": "Aldolasa",
    "1080": "Hepatitis B HBc Ac",
    "1086": "Hepatitis B HBs Ag",
    "1095": "Hepatitis C (HCV) Ac",
    "63": "HIV 4º Generación (Ag / Ac Combinado)",
    "371": "FTA - ABS",
    "9571": "Toxoplasmosis IgG",
    "9580": "Toxoplasmosis IgM captura",
    "665": "Micológico Cultivo Identificación",
    "190": "CPK (Creatinfosfoquinasa)",
    "594": "LDH (Láctico Dehidrogenasa)",
    "761": "Proteína C Reactiva (PCR) Cuantitativa",
    "1175": "Gonadotrofina Coriónica Beta (Beta hCG)",
    "370": "FSH (Hormona Foliculoestimulante)",
    "300": "Estradiol (E2)",
    "863.2": "Testosterona",
    "9375": "Testosterona Libre",
    "537": "Inmunoglobulina A (IgA)",
    "540": "Inmunoglobulina G (IgG)",
    "541": "Inmunoglobulina M (IgM)",
    "1115": "CA 125",
    "1000": "Antígeno Prostático Específico Total (PSA)",
    "343": "Ferremia (Hierro sérico)",
    "5230": "Ferritina",
    "875": "Transferrina",
    "837": "Transferrina Índice de Saturación",
    "9887": "Vitamina B6 (Piridoxina)",
    "345": "Fibrinógeno",
    "4418": "Dímero D",
    "759": "Prolactina",
    "612": "LH (Hormona Luteinizante)",
    "758": "Progesterona",
    "6": "ACTH (Adrenocorticotrofina)",
    "8580": "17 Hidroxi Progesterona sérica",
    "189": "Cortisol matutino",
    "189.1": "Cortisol vespertino",
    "189.Dexa": "Cortisol post Dexametasona",
    "4008": "Cortisol Libre Urinario (CLU)",
    "4012": "Cortisol saliva 23:00 hs",
    "739": "PTH (Parathormona)",
    "144": "CEA (Antígeno Carcinoembrionario)",
    "1125": "CA 19-9",
    "20": "Alfafetoproteína (AFP)",
    "3025": "Beta Cross Laps",
    "7939": "Osteocalcina",
    "6452": "Homocisteína basal",
    "6452.1": "Homocisteína post carga de Metionina",
    "1145": "Rubeola IgG",
    "1150": "Rubeola IgM",
    "1025": "Citomegalovirus (CMV) IgG",
    "1030": "Citomegalovirus (CMV) IgM",
    "6042": "Herpes Simplex 1 IgG",
    "6050": "Herpes Simplex 1 IgM",
    "6067": "Herpes Simplex 2 IgG",
    "6076": "Herpes Simplex 2 IgM",
    "243": "Chagas Ac (Electroquimioluminiscencia)",
    "242": "Chagas HAI",
    "243.1": "Chagas IgG IFI",
    "5888": "Hepatitis A (HAV) Ac Totales",
    "1075": "Hepatitis A (HAV) IgM",
    "9622.BF": "Transglutaminasa IgA",
    "9631": "Transglutaminasa IgG",
    "1000.L": "PSA Libre",
    "187": "Coprocultivo",
    "8218": "Beta 2 Glicoproteína Ac IgG",
    "8222": "Beta 2 Glicoproteína Ac IgM",
    "1082": "Hepatitis B HBc IgM",
}


def limpiar_dni(valor):
    if pd.isna(valor):
        return ""
    return str(valor).strip().replace(".", "").replace("-", "").split(",")[0].strip()


def limpiar_solicitud(valor):
    s = str(int(float(str(valor)))).strip()
    if s.startswith("9") and len(s) == 5:
        return s[1:]
    return s.zfill(4)


def procesar_gano(archivo_bytes):
    """Lee el Excel de Gano y devuelve lista de filas de derivaciones."""
    df = pd.read_excel(BytesIO(archivo_bytes), header=None)

    header_row = None
    for i, row in df.iterrows():
        if any(str(v).strip() == "Solicitud" for v in row):
            header_row = i
            break
    if header_row is None:
        raise ValueError("No se encontró fila de encabezado con 'Solicitud'")

    df.columns = df.iloc[header_row]
    df = df.iloc[header_row + 1:].reset_index(drop=True)
    df = df[df["Solicitud"].apply(lambda x: str(x).strip() not in ("nan", "") and
             not str(x).strip().startswith("Quimio"))].reset_index(drop=True)

    filas = []
    cols_presentes = set(df.columns)

    for _, row in df.iterrows():
        solicitud = row.get("Solicitud", "")
        if pd.isna(solicitud) or str(solicitud).strip() in ("nan", ""):
            continue

        nro_interno = limpiar_solicitud(solicitud)
        paciente = str(row.get("Paciente", "")).strip()
        dni = limpiar_dni(row.get("Documento", ""))

        for col_gano, codigo in DERIVAR.items():
            if col_gano not in cols_presentes:
                continue
            val = row.get(col_gano, None)
            if pd.isna(val):
                analisis = NOMBRE_ANALISIS.get(codigo, col_gano.strip())
                filas.append({
                    "BARCODE": "",
                    "NRO INTERNO": nro_interno,
                    "APELLIDO Y NOMBRE": paciente,
                    "DNI": dni,
                    "NRO CLIENTE": CODIGO_CLIENTE,
                    "CÓDIGO PRESTACION": codigo,
                    "FECHA ENVÍO": "",
                    "ANÁLISIS": analisis,
                })

    return filas


def generar_excel_derivaciones(filas):
    """Genera el Excel de derivaciones (para cargar barcodes) como bytes."""
    cols = ["BARCODE", "NRO INTERNO", "APELLIDO Y NOMBRE", "DNI",
            "NRO CLIENTE", "CÓDIGO PRESTACION", "FECHA ENVÍO", "ANÁLISIS"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Derivaciones"

    # Colores
    azul = PatternFill("solid", fgColor="1F4E79")
    celeste = PatternFill("solid", fgColor="BDD7EE")
    amarillo = PatternFill("solid", fgColor="FFFF99")

    # Fila 1: encabezados
    for c, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=col)
        cell.fill = azul
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")

    # Fila 2: referencia
    ref = ["STICKERS GORNITZ", "NRO PROTOCOLO", "", "", CODIGO_CLIENTE, "", "", ""]
    for c, val in enumerate(ref, 1):
        cell = ws.cell(row=2, column=c, value=val)
        cell.fill = celeste
        cell.font = Font(bold=True)

    # Datos
    for r, fila in enumerate(filas, 3):
        for c, col in enumerate(cols, 1):
            cell = ws.cell(row=r, column=c, value=str(fila[col]))
            cell.number_format = "@"
            # Columna BARCODE en amarillo para que se vea dónde pegar
            if c == 1:
                cell.fill = amarillo

    # Anchos
    anchos = [18, 14, 35, 12, 14, 18, 14, 40]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ancho

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generar_txt_gornitz(archivo_bytes):
    """Lee el Excel con barcodes y genera el TXT para Gornitz."""
    df = pd.read_excel(BytesIO(archivo_bytes), header=None)
    datos = df.iloc[2:].reset_index(drop=True)
    datos.columns = ["BARCODE", "NRO_INTERNO", "NOMBRE", "DNI",
                     "CLIENTE", "CODIGO", "FECHA", "ANALISIS"]

    lineas = []
    for _, row in datos.iterrows():
        barcode = str(row["BARCODE"]).strip()
        nro = str(row["NRO_INTERNO"]).strip()
        nombre = str(row["NOMBRE"]).strip()
        dni = str(row["DNI"]).strip()
        if dni.endswith(".0"):
            dni = dni[:-2]
        codigo = str(row["CODIGO"]).strip()
        if codigo.endswith(".0"):
            codigo = codigo[:-2]

        if barcode in ("nan", "") or nro in ("nan", "") or codigo in ("nan", ""):
            continue

        lineas.append(f"{barcode};{nro} {nombre};{CODIGO_CLIENTE};{codigo};V;;0;{dni};{dni}")

    return "\n".join(lineas), len(lineas)
